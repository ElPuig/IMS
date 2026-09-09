# -*- coding: utf-8 -*-

import base64
import csv
import io
import logging
from datetime import datetime

from markupsafe import Markup

from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class EmsStudentImportWizard(models.TransientModel):
    _name = "ems.student_import_wizard"
    _description = "Student import wizard (Esfera/SAGA xlsx)"

    # The only column the file must provide: it is what identifies the student, so it
    # doubles as the marker _find_headers scans for to locate a sheet's header row.
    _STUDENT_ID_COLUMN = "Identificador de l'alumne/a"
    # Import control rather than student data: a returning student must be reactivated
    # and restored to the 'student' contact type in both modes, whatever EMS holds now.
    _CONTROL_FIELDS = ('active', 'contact_type')

    file = fields.Binary(string="Esfera xlsx file", required=True)
    file_name = fields.Char()
    overwrite = fields.Boolean(
        string="Overwrite existing data",
        default=False,
        help="Leave it unticked to keep what EMS already holds: only fields that are "
             "currently empty get filled in, and students missing from EMS are created. "
             "Tick it to let the file's values replace EMS's ones. In both cases a column "
             "that comes empty in the file never erases an existing value, and notes are "
             "always kept.",
    )
    result_html = fields.Html(string="Import result", readonly=True)
    log_file = fields.Binary(string="Import log (CSV)", readonly=True)
    log_file_name = fields.Char()

    def action_import(self):
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("openpyxl is required to import xlsx files."))

        raw = base64.b64decode(self.file)
        wb = openpyxl.load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)

        stats = {'created': 0, 'updated': 0, 'errors': [], 'warnings': [], 'log': []}
        sheets_read = 0
        for ws in wb.worksheets:
            header_row_idx, col_map = self._find_headers(ws)
            if header_row_idx is None:
                # A sheet with no recognisable header row (an empty or auxiliary tab) is
                # skipped instead of aborting the whole file: an Esfera export regularly
                # splits its students across several sheets, and every one of them
                # holding data has to be imported.
                continue
            sheets_read += 1
            for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
                if not any(row):
                    continue
                try:
                    self._process_row(row, col_map, stats)
                except Exception as e:
                    _logger.warning("Error processing row: %s", e)
                    stats['errors'].append(str(e))

        if not sheets_read:
            raise UserError(_(
                "Could not find the header row in any sheet of the file. Make sure it "
                "contains a column '%(column)s'.",
                column=self._STUDENT_ID_COLUMN,
            ))

        self.log_file = self._build_log_csv(stats['log'])
        self.log_file_name = f"import_esfera_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        self.result_html = self._build_result_html(stats)
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'ems.student_import_wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _find_headers(self, ws):
        def normalize(s):
            return str(s or '').strip().replace('’', "'").replace('‘', "'")
        for idx, row in enumerate(ws.iter_rows(max_row=20, values_only=True), start=1):
            if row and any(normalize(c) == self._STUDENT_ID_COLUMN for c in row):
                col_map = {normalize(c): i for i, c in enumerate(row) if c}
                return idx, col_map
        return None, {}

    @staticmethod
    def _is_empty(value):
        """Whether a value counts as "nothing to write".

        Covers what both sides of the write policy can hold: None/False from an
        absent column, a blank or whitespace-only cell, and an empty recordset
        read back from a many2one field.
        """
        if isinstance(value, str):
            return not value.strip()
        if value is None:
            return True
        return not value

    def _values_to_write(self, record, vals):
        """Filter the file's values down to what may actually be written on an
        existing record. Two rules, applied to students and family contacts alike:

        - a column that came empty in the file never blanks a value EMS already
          holds (in either mode);
        - without ``overwrite``, EMS's own data always wins, so only fields that
          are currently empty get filled in.
        """
        to_write = {}
        for field_name, value in vals.items():
            if self._is_empty(value):
                continue
            if field_name not in self._CONTROL_FIELDS and not self.overwrite and not self._is_empty(record[field_name]):
                continue
            to_write[field_name] = value
        return to_write

    def _prepend_import_notes(self, record, notes):
        """Stack this import's notes on top of the record's existing ones.

        Notes are never overwritten, in either mode: the new block is stamped with
        the import date and closed by a horizontal rule, so it stays visible when
        each block arrived and nothing written by hand is ever lost.
        """
        if not notes:
            return
        stamp = fields.Datetime.context_timestamp(record, fields.Datetime.now()).strftime('%d/%m/%Y %H:%M')
        block = Markup('<p><strong>{stamp}</strong></p>{notes}<hr/>').format(
            stamp=stamp, notes=Markup(notes),
        )
        record.comment = block + (record.comment or '')

    def _col_get(self, row, col_map, col_name):
        """Read one Esfera column by its exact header name, or None if the
        column is absent from this file or the cell is empty for this row.
        Shared by _process_row (student columns) and _process_tutor (Tutor
        N - prefixed columns)."""
        idx = col_map.get(col_name)
        if idx is None:
            return None
        val = row[idx] if idx < len(row) else None
        return str(val).strip() if val is not None else None

    def _process_row(self, row, col_map, stats):
        def get(col_name):
            return self._col_get(row, col_map, col_name)

        # The student identifier is the only mandatory column: it is what matches the
        # row to a student, so a row without one can be neither found nor created.
        ralc = get(self._STUDENT_ID_COLUMN)
        if not ralc:
            stats['warnings'].append(_(
                "A row was skipped: it carries no student identifier (%(column)s).",
                column=self._STUDENT_ID_COLUMN,
            ))
            return

        # Group - searched by external_id (Esfera code) whenever the file provides one.
        # Normalize whitespace: Esfera sometimes uses multiple spaces (e.g. "CFPM    IC10201")
        esfera_code = ' '.join((get('Grup Classe') or '').split()) or None
        group = self.env['ems.group'].search(
            [('external_id', '=', esfera_code)], limit=1) if esfera_code else self.env['ems.group']

        # Student name
        firstname = get('Nom') or ''
        surname1 = get('Primer Cognom') or ''
        surname2 = get('Segon Cognom') or ''
        name = ' '.join(filter(None, [firstname, surname1, surname2]))

        # Documents
        doc_nums = get('Número de document d\'identitat')
        doc_types = get('Tipus de document d\'identitat') or ''
        docs = self._parse_documents(doc_nums, doc_types)

        # Personal data
        birth_str = get('Data naixement')
        birth_date = self._parse_date(birth_str)
        citizenship = self._find_country(get('Nacionalitat'))
        birth_country = self._find_country(get('País naixement'))
        raw_phone = get('Telèfon') or get('Contacte alumne - Telèfon')
        student_phone, student_mobile = self._split_phone_mobile(raw_phone)
        email = get('Correu electrònic') or get('Contacte alumne - Correu electrònic')

        # Address
        street = self._build_street(
            get('Tipus de via'), get('Nom via'), get('Número'),
            get('Bloc'), get('Escala'), get('Planta'), get('Porta'),
            get('Resta de dades de l\'adreça')
        )
        city = get('Municipi de residència')
        zip_code = get('Codi postal')
        country = self._find_country(get('País de residència'))
        state = self._find_state(get('Província de residència'), country.id if country else False)

        # Extra notes
        notes = self._build_student_notes(get, esfera_code, group)
        if esfera_code and not group:
            # Intentional (not every group may exist in EMS yet at import time), but
            # surfaced in the result summary too, not just the comment note above —
            # see plans/student_import_wizard_data_quality_gaps.md (now resolved).
            stats['warnings'].append(_(
                "%(name)s: no group found for Esfera code '%(code)s' — imported "
                "without a group, needs manual placement.",
                name=name, code=esfera_code,
            ))

        student_data = {
            'name': name,
            'contact_type': 'student',
            'birth_date': birth_date,
            'document_id': docs.get('DNI') or docs.get('NIE'),
            'passport_id': docs.get('PASS') or docs.get('Passaport'),
            'medical_id': docs.get('TIS'),
            'phone': student_phone,
            'mobile': student_mobile,
            'email': email,
            'street': street,
            'city': city,
            'zip': zip_code,
            'country_id': country.id if country else False,
            'state_id': state.id if state else False,
            'citizenship_id': citizenship.id if citizenship else False,
            'birth_country_id': birth_country.id if birth_country else False,
            'main_group_id': group.id if group else False,
            'student_id': ralc,
            # Re-admits an ex-student (alumni/withdrawal) archived on exit: without
            # this, an existing-but-inactive match is written but stays archived.
            'active': True,
        }

        student = self._get_or_create_student(ralc, student_data, stats, notes=notes)
        if not student:
            return

        # Tutors
        for prefix in ['Tutor 1', 'Tutor 2']:
            self._process_tutor(row, col_map, prefix, student, stats)

    # NOTE: the note labels below are deliberately kept in Catalan, untranslated,
    # even though they end up in a user-visible Notes field. They are a verbatim
    # echo of Esfera/SAGA's own official Catalan field names (the Catalan
    # education administration's system of record) — translating them would
    # weaken traceability back to "this is exactly what Esfera exported for this
    # student", which is the whole point of keeping them. See student_import_wizard.md.
    def _build_student_notes(self, get, esfera_code, group):
        lines = []
        if esfera_code and not group:
            lines.append(f"Grup Classe (SAGA): {esfera_code}")
        for label, key in [
            ('Província de naixement', 'Província naixement'),
            ('Municipi de naixement', 'Municipi naixement'),
            ('Localitat de residència', 'Localitat de residència'),
            ('Alumne tutelat legalment', 'Alumne tutelat legalment'),
            ('Alumne emancipat legalment', 'Alumne emancipat legalment'),
            ('Alumne amb custòdia compartida en dos domicilis', 'Alumne amb custòdia compartida en dos domicilis'),
            ('Contacte altres - Tipus', 'Contacte altres alumne - Tipus'),
            ('Contacte altres - Valor', 'Contacte altres alumne - Valor'),
            ('Contacte altres - Observacions', 'Contacte altres alumne - Observacions'),
            ('Contacte propis - Observacions', 'Contacte propis alumne - Observacions'),
            ('Observacions', 'Observacions'),
        ]:
            val = get(key)
            if val and val.lower() not in ('no', 'false', ''):
                lines.append(f"{label}: {val}")
        return '<br/>'.join(lines) if lines else False

    def _get_or_create_student(self, ralc, data, stats, notes=None):
        existing = False
        if ralc:
            existing = self.env['res.partner'].with_context(active_test=False).search(
                [('student_id', '=', ralc)], limit=1)
        if existing:
            to_write = self._values_to_write(existing, data)
            if to_write:
                existing.write(to_write)
            self._prepend_import_notes(existing, notes)
            stats['updated'] += 1
            stats['log'].append({'tipus': 'Alumne', 'accio': 'Actualitzat', 'partner_id': existing.id, 'ts': datetime.now()})
            return existing
        vals = {name: value for name, value in data.items() if not self._is_empty(value)}
        if not vals.get('name'):
            # res.partner.name is required, and a nameless contact would be unusable
            # anyway — an existing student is unaffected, since their name is never read.
            stats['warnings'].append(_(
                "Student '%(ralc)s' was skipped: the row carries no name and no student "
                "with that identifier exists in EMS yet.",
                ralc=ralc,
            ))
            return False
        student = self.env['res.partner'].create(vals)
        self._prepend_import_notes(student, notes)
        stats['created'] += 1
        stats['log'].append({'tipus': 'Alumne', 'accio': 'Creat', 'partner_id': student.id, 'ts': datetime.now()})
        return student

    def _process_tutor(self, row, col_map, prefix, student, stats):
        def get(col_name):
            return self._col_get(row, col_map, col_name)

        nom = get(f'{prefix} - nom')
        if not nom:
            return

        # Note: Esfera exports '1r cognom ' with trailing space
        surname1 = get(f'{prefix} - 1r cognom ') or get(f'{prefix} - 1r cognom') or ''
        surname2 = get(f'{prefix} - 2n cognom') or ''
        full_name = ' '.join(filter(None, [nom, surname1, surname2]))

        doc_num = get(f'{prefix} - doc. identitat')
        docs = self._parse_documents(doc_num, '')

        tutor_num = '1er' if '1' in prefix else '2on'
        contact_raw = get(f'Contacte {tutor_num} tutor alumne - Valor')
        raw_phone, email = self._parse_contact_value(contact_raw)
        phone, mobile = self._split_phone_mobile(raw_phone)
        observacio = get(f'Contacte {tutor_num} tutor alumne - Observacions')

        street = self._build_street(
            get(f'{prefix} - tipus via'), get(f'{prefix} - nom via'),
            get(f'{prefix} - número'), get(f'{prefix} - bloc'),
            get(f'{prefix} - escala'), get(f'{prefix} - planta'),
            get(f'{prefix} - porta'), get(f'{prefix} - resta de dades de l\'adreça')
        )
        city = get(f'{prefix} - municipi')
        zip_code = get(f'{prefix} - CP')
        country = self._find_country(get(f'{prefix} - país'))
        state = self._find_state(get(f'{prefix} - provincia'), country.id if country else False)

        # Extra notes for tutor — same Catalan-verbatim rationale as _build_student_notes.
        tutor_notes = []
        for label, key in [
            ('Localitat', f'{prefix} - localitat'),
            ('Destinatari correspondència', f'{prefix} - destinatari correspondència'),
            ('Persona jurídica', f'{prefix} - persona jurídica'),
            ('Rep notificacions', f'{prefix} - contactes: rebre notificacions'),
        ]:
            val = get(key)
            if val and val.lower() not in ('no', 'false', ''):
                tutor_notes.append(f"{label}: {val}")
        if prefix == 'Tutor 2':
            shared = get('Tutors comparteixen domicili')
            if shared and shared.lower() not in ('no', 'false', ''):
                tutor_notes.append(f"Comparteix domicili amb Tutor 1: {shared}")

        family, accio = self._get_or_create_family(
            full_name, doc_num, phone, mobile, email,
            {'street': street, 'city': city, 'zip': zip_code,
             'country_id': country.id if country else False,
             'state_id': state.id if state else False},
            notes='<br/>'.join(tutor_notes) if tutor_notes else False,
        )
        if not family:
            return
        if not doc_num:
            # KNOWN LIMITATION, kept intentionally (see
            # plans/student_import_wizard_data_quality_gaps.md, now resolved):
            # dedup only matches on document number, so a documentless tutor always
            # creates a new family contact. A fuzzier name/phone fallback was
            # rejected (false-positive merge risk) - surfaced as a warning instead,
            # so it's visible in the result summary for manual review.
            stats['warnings'].append(_(
                "%(student)s: tutor '%(tutor)s' has no document number — dedup "
                "skipped, a new family contact may have been created even if one "
                "already exists.",
                student=student.name, tutor=full_name,
            ))
        stats['log'].append({'tipus': 'Familiar', 'accio': accio, 'partner_id': family.id, 'ts': datetime.now()})

        relation_type, is_fallback = self._deduce_relation_type(observacio)
        if is_fallback and observacio:
            note = f"[Import Esfera] Relació {prefix}: '{observacio}' (assignada com a Tutor per defecte)"
            student.comment = f"{student.comment}<br/>{note}".strip() if student.comment else note

        self._link_family_to_student(family, student, relation_type)

    def _get_or_create_family(self, name, doc_num, phone, mobile, email, address_data, notes=None):
        """Find or create the family contact for a tutor row.

        Values are filtered through _values_to_write, so a family contact follows
        exactly the same policy as the student: empty cells never blank anything,
        and EMS's data wins unless 'overwrite' is ticked.

        KNOWN LIMITATION, kept intentionally (see student_import_wizard.md): dedup
        only matches on doc_num (document_id/passport_id). A tutor row with no
        document number always creates a new family partner — there is no
        name/phone/email fallback match (rejected: false-positive merge risk is
        worse than a duplicate contact). The caller (_process_tutor) surfaces this
        in stats['warnings'] instead, so it's visible for manual review.
        """
        if not name:
            return False, None

        domain = [('contact_type', '=', 'family')]
        existing = False
        if doc_num:
            existing = (
                self.env['res.partner'].search(domain + [('document_id', '=', doc_num)], limit=1)
                or self.env['res.partner'].search(domain + [('passport_id', '=', doc_num)], limit=1)
            )
        family_vals = dict(address_data, **{
            'name': name,
            'contact_type': 'family',
            'document_id': doc_num,
            'phone': phone,
            'mobile': mobile,
            'email': email,
        })
        if existing:
            to_write = self._values_to_write(existing, family_vals)
            if to_write:
                existing.write(to_write)
            self._prepend_import_notes(existing, notes)
            return existing, 'Actualitzat'

        vals = {field: value for field, value in family_vals.items() if not self._is_empty(value)}
        family = self.env['res.partner'].create(vals)
        self._prepend_import_notes(family, notes)
        return family, 'Creat'

    def _link_family_to_student(self, family, student, relation_type):
        existing = self.env['res.partner.relation'].search([
            ('left_partner_id', '=', family.id),
            ('right_partner_id', '=', student.id),
        ], limit=1)
        if not existing:
            self.env['res.partner.relation'].create({
                'left_partner_id': family.id,
                'type_id': relation_type.id,
                'right_partner_id': student.id,
            })

    def _deduce_relation_type(self, text):
        t = (text or '').lower()
        is_fallback = False
        if 'mare' in t or 'madre' in t:
            rel = self.env.ref('ems.relation_type_mother')
        elif 'pare' in t or 'padre' in t:
            rel = self.env.ref('ems.relation_type_father')
        elif 'àvia' in t or 'avia' in t or 'abuela' in t:
            rel = self.env.ref('ems.relation_type_grandmother')
        elif 'avi' in t or 'abuelo' in t:
            rel = self.env.ref('ems.relation_type_grandfather')
        elif 'tieta' in t or 'tia' in t or 'oncle' in t or 'tio' in t:
            rel = self.env.ref('ems.relation_type_uncle_aunt')
        elif 'germana' in t or 'hermana' in t or 'germà' in t or 'hermano' in t:
            rel = self.env.ref('ems.relation_type_sibling')
        else:
            rel = self.env.ref('ems.relation_type_tutor')
            is_fallback = True
        return rel, is_fallback

    def _parse_documents(self, numbers_str, types_str):
        if not numbers_str:
            return {}
        numbers = [n.strip() for n in numbers_str.split(' - ')]
        types = [t.strip() for t in (types_str or '').split(' - ')]
        return dict(zip(types, numbers))

    def _parse_date(self, value):
        if not value:
            return False
        for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
            try:
                return datetime.strptime(str(value), fmt).date()
            except ValueError:
                continue
        return False

    def _build_street(self, tipus_via, nom_via, numero, bloc, escala, planta, porta, resta):
        parts = [tipus_via, nom_via, numero, bloc, escala, planta, porta, resta]
        return ' '.join(p for p in parts if p and str(p).strip()) or False

    def _parse_contact_value(self, raw):
        if not raw:
            return None, None
        parts = [p.strip() for p in raw.split(' - ')]
        phone = parts[0] if parts else None
        email = parts[1] if len(parts) > 1 else None
        return phone, email

    def _split_phone_mobile(self, number, country_code='ES'):
        if not number:
            return None, None
        try:
            import phonenumbers
            parsed = phonenumbers.parse(number, country_code)
            if phonenumbers.number_type(parsed) == phonenumbers.PhoneNumberType.MOBILE:
                return None, number
            return number, None
        except Exception:
            return number, None

    def _find_country(self, name):
        if not name:
            return False
        return self.env['res.country'].with_context(lang='ca_ES').search(
            [('name', 'ilike', name)], limit=1
        )

    def _find_state(self, name, country_id):
        if not name or not country_id:
            return False
        return self.env['res.country.state'].with_context(lang='ca_ES').search(
            [('name', 'ilike', name), ('country_id', '=', country_id)], limit=1
        )

    def _build_log_csv(self, log_entries):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            'tipus', 'accio', 'data_hora', 'id', 'nom', 'document_id',
            'email', 'telèfon', 'mòbil', 'grup', 'vinculat_a', 'vinculat_als_ids',
        ])
        for entry in log_entries:
            partner = self.env['res.partner'].browse(entry['partner_id'])
            if entry['tipus'] == 'Alumne':
                relations = self.env['res.partner.relation'].search([('right_partner_id', '=', partner.id)])
                linked = relations.mapped('left_partner_id')
                grup = partner.main_group_id.display_name if partner.main_group_id else ''
            else:
                relations = self.env['res.partner.relation'].search([('left_partner_id', '=', partner.id)])
                linked = relations.mapped('right_partner_id')
                grup = ''
            writer.writerow([
                entry['tipus'],
                entry['accio'],
                entry['ts'].strftime('%Y-%m-%d %H:%M:%S'),
                partner.id,
                partner.name,
                partner.document_id or '',
                partner.email or '',
                partner.phone or '',
                partner.mobile or '',
                grup,
                ', '.join(linked.mapped('name')),
                ', '.join(str(p.id) for p in linked),
            ])
        return base64.b64encode(output.getvalue().encode('utf-8-sig')).decode()

    def _build_result_html(self, stats):
        errors_html = ''
        if stats['errors']:
            errors_html = Markup('<p><strong>{}</strong></p>{}').format(
                _("Errors (%(count)s):", count=len(stats['errors'])),
                self.env['ems.base'].build_html_list(stats['errors']),
            )
        warnings = stats.get('warnings', [])
        warnings_html = ''
        if warnings:
            warnings_html = Markup('<p><strong>{}</strong></p>{}').format(
                _("Warnings (%(count)s):", count=len(warnings)),
                self.env['ems.base'].build_html_list(warnings),
            )
        return Markup(
            '<p>✅ <strong>{created_label}</strong> {created}</p>'
            '<p>🔄 <strong>{updated_label}</strong> {updated}</p>'
            '{warnings_html}'
            '{errors_html}'
        ).format(
            created_label=_("Students created:"), created=stats['created'],
            updated_label=_("Students updated:"), updated=stats['updated'],
            warnings_html=warnings_html,
            errors_html=errors_html,
        )
